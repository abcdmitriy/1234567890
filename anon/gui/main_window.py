# gui/main_window.py
import html
import os
import json
import re
from datetime import datetime
from PyQt5.QtWidgets import (QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
                             QPushButton, QSplitter, QLabel, QListWidget,
                             QFileDialog, QMessageBox, QTableWidget,
                             QTableWidgetItem, QProgressBar, QStatusBar,
                             QHeaderView, QTabWidget, QTextEdit, QMenuBar, QDialog,
                             QTextBrowser, QFrame, QComboBox, QCheckBox,
                             QListWidgetItem, QTabBar, QTreeWidget, QTreeWidgetItem,
                             QMenu, QLineEdit)
from PyQt5.QtCore import Qt, QEvent
from PyQt5.QtGui import QColor, QTextCursor

from core.detectors.detector_manager import DetectorManager
from core.parsers.docx_parser import DOCXParser
from core.parsers.pdf_parser import PDFParser
from core.parsers.image_parser import ImageParser
from core.processors.deanon import Deanon
from storage.map_manager import MapManager
from .settings_dialog import SettingsDialog
from .anonymize_worker import AnonymizeWorker
from .deanon_window import DeanonWindow


class MainWindow(QMainWindow):
    VIRTUAL_TEXT_PATH = "<virtual:my_text>"
    
    def __init__(self):
        super().__init__()

        self.setWindowTitle("Анонимизатор документов")
        self.setGeometry(100, 100, 1400, 900)
        self.setStyleSheet(self._build_stylesheet())

        self.settings = {
            "fio": True,
            "passport_series": True,
            "passport_code": True,
            "addresses": True,
            "phones": True,
            "emails": True,
            "org": True,
            "inn": True,
            "bank_account": True,
            "bic": True
        }
        self.save_folder = None
        self.preserve_structure = True
        self.selected_files = []
        self.file_relative_paths = {}
        self.worker = None
        self._suppress_processing_notifications = False
        self.current_mapping = {}
        self.base_mapping_snapshot = {}
        self.excluded_aliases = set()
        self.selected_mapping_alias = None
        self.current_file_replacements = {}
        self.preview_text_cache = {}
        self.file_text_cache = {}
        self.original_text_cache = {}
        self.anonymized_file_texts = {}
        self.active_file_path = None
        self.last_anonymized_result = None
        self.deanon_reverse_map = {}
        self.map_manager = MapManager()
        self.available_maps = []
        self.current_map_name = None
        self.current_map_path = None
        self.setAcceptDrops(True)

        self.left_tabs = QTabWidget()
        self.files_tab = QWidget()
        self.mapping_tab = QWidget()

        self.create_left_tabs_ui()
        self.create_main_layout()
        self.setup_menu_bar()
        self.refresh_controls()

    def _build_stylesheet(self):
        return """
        QMainWindow, QDialog {
            background: #f5f7fb;
            color: #111827;
            font-family: 'Segoe UI', sans-serif;
        }
        QWidget {
            font-family: 'Segoe UI', sans-serif;
        }
        QPushButton {
            background: #ffffff;
            color: #111827;
            border: 1px solid #dfe3ec;
            border-radius: 10px;
            padding: 8px 14px;
            min-height: 32px;
            font-weight: 600;
        }
        QPushButton:hover {
            background: #f3f6ff;
            border-color: #bfd1ff;
        }
        QPushButton:pressed {
            background: #eaf0ff;
        }
        QPushButton:disabled {
            background: #eef2f8;
            color: #9aa4b2;
            border-color: #e5e7eb;
        }
        QListWidget, QTreeWidget, QTableWidget, QTextEdit, QComboBox, QLineEdit {
            border: 1px solid #e2e8f0;
            border-radius: 12px;
            background: #ffffff;
            padding: 8px;
            color: #111827;
            selection-background-color: #e0e7ff;
            selection-color: #111827;
        }
        QTreeWidget {
            outline: none;
        }
        QTreeWidget::item {
            padding: 6px 8px;
            border-radius: 8px;
            margin: 2px 0;
        }
        QTreeWidget::item:selected {
            background: #e0e7ff;
            color: #111827;
        }
        QTableWidget {
            gridline-color: #edf2f7;
        }
        QHeaderView::section {
            background: #f8fafc;
            color: #475569;
            padding: 8px;
            border: none;
            font-weight: 700;
        }
        QTextBrowser {
            background: #ffffff;
            border: 1px solid #e2e8f0;
            border-radius: 12px;
            padding: 12px;
            color: #111827;
        }
        QProgressBar {
            border-radius: 8px;
            border: 1px solid #d9e2f1;
            background: #edf2f7;
            text-align: center;
        }
        QProgressBar::chunk {
            background: qlineargradient(x1:0, y1:0, x2:1, y2:0, stop:0 #4f46e5, stop:1 #8b5cf6);
            border-radius: 7px;
        }
        QLabel {
            color: #374151;
        }
        QStatusBar {
            background: #eff3f9;
            color: #111827;
        }
        QTabWidget::pane {
            border: 1px solid #e2e8f0;
            border-radius: 12px;
            background: #f8fafc;
        }
        QTabBar::tab {
            background: #edf2f7;
            border: 1px solid #e2e8f0;
            border-bottom: none;
            border-top-left-radius: 10px;
            border-top-right-radius: 10px;
            min-width: 180px;
            min-height: 30px;
            padding: 4px 18px;
            color: #475569;
            font-weight: 600;
        }
        QTabBar::tab:selected {
            background: #ffffff;
            color: #111827;
            font-weight: 700;
        }
        QCheckBox {
            color: #334155;
            spacing: 10px;
            padding: 6px 0;
        }
        QGroupBox {
            border: 1px solid #e2e8f0;
            border-radius: 12px;
            margin-top: 12px;
            background: #ffffff;
            color: #1f2937;
            font-weight: 600;
            padding-top: 10px;
        }
        QGroupBox::title {
            subcontrol-origin: margin;
            left: 14px;
            padding: 0 8px;
            color: #475569;
        }
        QComboBox {
            padding-left: 10px;
            min-height: 32px;
        }
        QComboBox::drop-down {
            subcontrol-origin: padding;
            subcontrol-position: top right;
            border-left: 1px solid #e2e8f0;
            width: 28px;
            background: #f8fafc;
        }
        QMenu {
            background: #ffffff;
            border: 1px solid #e2e8f0;
            border-radius: 10px;
            color: #111827;
        }
        """

    def create_left_tabs_ui(self):
        self.files_tab = QWidget()
        file_upload_layout = QVBoxLayout(self.files_tab)
        file_upload_layout.setAlignment(Qt.AlignTop)

        self.file_list_widget = QListWidget()
        self.file_list_widget.setSelectionMode(QListWidget.SingleSelection)
        self.file_list_widget.itemClicked.connect(self.on_file_selected)
        self.file_list_widget.setVisible(False)
        file_upload_layout.addWidget(self.file_list_widget, stretch=2)

        self.info_label = QLabel("Загружено файлов: 0")
        self.info_label.setAlignment(Qt.AlignLeft)
        file_upload_layout.addWidget(self.info_label)

        mapping_layout = QVBoxLayout(self.mapping_tab)
        mapping_layout.setAlignment(Qt.AlignTop)

        self.mapping_table = QTableWidget()
        self.mapping_table.setColumnCount(5)
        self.mapping_table.setHorizontalHeaderLabels([
            "Псевдоним", "Исходное значение", "Категория",
            "Файлы", "Действие"
        ])
        self.mapping_table.setColumnHidden(2, True)
        self.mapping_table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.mapping_table.setEditTriggers(QTableWidget.EditKeyPressed)
        self.mapping_table.setVerticalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
        self.mapping_table.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
        self.mapping_table.itemSelectionChanged.connect(self.on_mapping_selection_changed)
        self.mapping_table.horizontalHeader().sectionClicked.connect(self.on_mapping_header_clicked)

        mapping_layout.addWidget(self.mapping_table, stretch=3)

        self.map_combo = QComboBox()
        self.map_combo.setEditable(False)
        self.map_combo.currentIndexChanged.connect(self.on_map_combo_changed)

        self.left_tabs.addTab(self.mapping_tab, "Карта соответствий")
        self.left_tabs.tabBar().hide()
        self.refresh_map_list()

    def create_main_layout(self):
        central_widget = QWidget()
        self.setCentralWidget(central_widget)

        main_splitter = QSplitter(Qt.Horizontal)
        central_widget.setLayout(QVBoxLayout())
        central_widget.layout().addWidget(main_splitter)

        left_panel = QWidget()
        main_left_layout = QVBoxLayout(left_panel)

        actions_panel = QFrame()
        actions_panel.setObjectName("panel")
        actions_layout = QVBoxLayout(actions_panel)
        actions_layout.setSpacing(10)

        first_row = QHBoxLayout()
        self.btn_settings = QPushButton("⚙")
        self.btn_settings.setToolTip("Настройки")
        self.btn_settings.setFixedWidth(54)
        self.btn_settings.clicked.connect(self.show_settings)

        self.btn_anonymize = QPushButton("Анонимизировать")
        self.btn_anonymize.clicked.connect(self.start_anonymization)
        first_row.addWidget(self.btn_settings)
        first_row.addWidget(self.btn_anonymize)
        actions_layout.addLayout(first_row)

        self.btn_deanonymize = QPushButton("Деанонимизировать")
        self.btn_deanonymize.clicked.connect(self.deanonymize_window)
        actions_layout.addWidget(self.btn_deanonymize)

        second_row = QHBoxLayout()
        self.btn_structure = QPushButton("Структура")
        self.btn_load = QPushButton("Загрузить")
        self.btn_structure.clicked.connect(self.show_structure_dialog)
        self.btn_load.clicked.connect(self.show_load_dialog)
        second_row.addWidget(self.btn_structure)
        second_row.addWidget(self.btn_load)
        actions_layout.addLayout(second_row)

        self.btn_save_results = QPushButton("Сохранить")
        self.btn_save_results.clicked.connect(self.show_save_dialog)
        actions_layout.addWidget(self.btn_save_results)

        self.btn_clear = QPushButton("Очистить")
        self.btn_clear.clicked.connect(self.clear_selected_files)
        actions_layout.addWidget(self.btn_clear)

        self.save_all_checkbox = QCheckBox("Сохранить все документы")
        self.save_all_checkbox.setChecked(False)
        self.save_all_checkbox.setVisible(False)

        main_left_layout.addWidget(actions_panel, stretch=1)
        main_left_layout.addWidget(self.left_tabs, stretch=2)

        main_splitter.addWidget(left_panel)

        right_container = QWidget()
        right_layout = QVBoxLayout(right_container)
        right_layout.setContentsMargins(0, 0, 0, 0)
        right_layout.setSpacing(0)

        self.file_tab_bar = QTabBar()
        self.file_tab_bar.setExpanding(False)
        self.file_tab_bar.setDocumentMode(True)
        self.file_tab_bar.setDrawBase(False)
        self.file_tab_bar.setTabsClosable(True)
        self.file_tab_bar.setAutoHide(False)
        self.file_tab_bar.currentChanged.connect(self.on_file_tab_changed)
        self.file_tab_bar.tabCloseRequested.connect(self.on_file_tab_close_requested)
        right_layout.addWidget(self.file_tab_bar)

        self.text_input_editor = QTextEdit()
        self.text_input_editor.setMinimumHeight(180)
        self.text_input_editor.setReadOnly(False)
        self.text_input_editor.setAcceptDrops(False)
        self.text_input_editor.setContentsMargins(0, 0, 0, 0)
        self.text_input_editor.setPlaceholderText("Вставьте или начните печатать")
        self.text_input_editor.textChanged.connect(self.on_editor_text_changed)
        self.text_input_editor.installEventFilter(self)
        self.text_input_editor.textChanged.connect(self.refresh_controls)
        self.text_input_editor.setStyleSheet(
            "QTextEdit { color: #111827 !important; background: #ffffff; }"
        )
        self.text_input_editor.setContextMenuPolicy(Qt.CustomContextMenu)
        self.text_input_editor.customContextMenuRequested.connect(self.on_editor_context_menu)
        right_layout.addWidget(self.text_input_editor, stretch=10)

        self.preview_browser = QTextBrowser()
        self.preview_browser.setReadOnly(True)
        self.preview_browser.setVisible(False)
        self.preview_browser.setHtml("<p style='color:#64748b;'>Подсветка замен будет показана здесь после анонимизации.</p>")
        right_layout.addWidget(self.preview_browser, stretch=0)

        main_splitter.addWidget(right_container)
        main_splitter.setStretchFactor(0, 35)
        main_splitter.setStretchFactor(1, 65)

        self.statusbar = QStatusBar()
        self.setStatusBar(self.statusbar)
        self.statusbar.showMessage("Готов к работе. Загрузите файлы для обработки.")

    def setup_menu_bar(self):
        menubar = QMenuBar()
        file_menu = menubar.addMenu("Файл")

        exit_action = file_menu.addAction("Выход")
        exit_action.triggered.connect(self.close)
        self.setMenuBar(menubar)

    def show_load_dialog(self):
        dialog = QDialog(self)
        dialog.setWindowTitle("Загрузить файлы или папку")
        dialog.resize(320, 140)
        dialog.setStyleSheet(self._build_stylesheet())
        layout = QVBoxLayout(dialog)

        btn_files = QPushButton("Выбрать файлы")
        btn_folder = QPushButton("Выбрать папку")

        def load_files():
            file_paths, _ = QFileDialog.getOpenFileNames(
                self,
                "Выберите файлы",
                "",
                "Все поддерживаемые файлы (*.docx *.pdf *.jpg *.jpeg *.png)"
            )
            if file_paths:
                for file_path in file_paths:
                    self.add_file_to_list(file_path, os.path.basename(file_path))
                last_path = file_paths[-1]
                self.show_file_preview(last_path)
                self.populate_text_editor_from_file(last_path)
            dialog.close()

        def load_folder():
            folder_path = QFileDialog.getExistingDirectory(self, "Выберите папку с файлами")
            if folder_path:
                for root, _, files in os.walk(folder_path):
                    for file in files:
                        if file.endswith(('.docx', '.pdf', '.jpg', '.jpeg', '.png')):
                            full_path = os.path.join(root, file)
                            relative_path = os.path.relpath(full_path, folder_path)
                            self.add_file_to_list(full_path, relative_path)
                if self.selected_files:
                    last_path = self.selected_files[-1]
                    self.show_file_preview(last_path)
                    self.populate_text_editor_from_file(last_path)
            dialog.close()

        btn_files.clicked.connect(load_files)
        btn_folder.clicked.connect(load_folder)
        layout.addWidget(btn_files)
        layout.addWidget(btn_folder)
        dialog.exec_()

    def show_structure_dialog(self):
        if not self.selected_files:
            QMessageBox.information(self, "Структура", "Нет загруженных файлов.")
            return

        dialog = QDialog(self)
        dialog.setWindowTitle("Структура загруженных файлов")
        dialog.resize(500, 400)
        dialog.setStyleSheet(self._build_stylesheet())
        layout = QVBoxLayout(dialog)

        tree_widget = QTreeWidget()
        tree_widget.setHeaderHidden(True)
        tree_widget.setColumnCount(1)
        tree_widget.setAlternatingRowColors(False)
        tree_widget.setUniformRowHeights(True)
        tree_widget.setRootIsDecorated(True)

        root_item = QTreeWidgetItem(tree_widget, ["Файлы"])
        root_item.setExpanded(True)
        root_item.setFlags(root_item.flags() & ~Qt.ItemIsUserCheckable)
        
        folder_map = {}
        leaf_items = []
        for file_path in self.selected_files:
            if file_path == self.VIRTUAL_TEXT_PATH:
                leaf_name = "my text"
                leaf_item = QTreeWidgetItem(root_item, [leaf_name])
            else:
                rel_name = self.file_relative_paths.get(file_path, os.path.basename(file_path)).replace('\\', '/')
                parts = [p for p in rel_name.split('/') if p]
                parent_item = root_item
                for idx, part in enumerate(parts[:-1]):
                    current_key = '/'.join(parts[:idx + 1])
                    child = folder_map.get(current_key)
                    if child is None:
                        child = QTreeWidgetItem(parent_item, [part])
                        child.setExpanded(True)
                        child.setFlags(child.flags() & ~Qt.ItemIsUserCheckable)
                        folder_map[current_key] = child
                    parent_item = child

                leaf_name = parts[-1] if parts else os.path.basename(file_path)
                leaf_item = QTreeWidgetItem(parent_item, [leaf_name])
                
            leaf_item.setData(0, Qt.UserRole, file_path)
            leaf_item.setData(0, Qt.UserRole + 1, self._file_display_name(file_path))
            leaf_items.append((leaf_item, file_path))
            
            if file_path == self.active_file_path:
                leaf_item.setBackground(0, QColor("#e0e7ff"))
                leaf_item.setForeground(0, QColor("#111827"))

        def on_item_double_clicked(item, column):
            file_path = item.data(0, Qt.UserRole)
            if file_path and file_path in self.selected_files:
                for leaf_item, path in leaf_items:
                    if path == self.active_file_path and path != file_path:
                        leaf_item.setBackground(0, QColor("white"))
                    elif path == file_path:
                        leaf_item.setBackground(0, QColor("#e0e7ff"))
                
                self.active_file_path = file_path
                self.selected_mapping_alias = None
                self.mapping_table.blockSignals(True)
                self.mapping_table.clearSelection()
                self.mapping_table.setCurrentCell(-1, -1)
                self.mapping_table.blockSignals(False)
                
                if self.file_list_widget.count() > self.selected_files.index(file_path):
                    self.file_list_widget.setCurrentRow(self.selected_files.index(file_path))
                self.show_file_preview(file_path)
                self._render_current_active_file_in_editor()
                self.update_file_count()
                self.statusbar.showMessage(f"Выбран файл: {self._file_display_name(file_path)}")
                self.refresh_controls()
                dialog.close()

        tree_widget.itemDoubleClicked.connect(on_item_double_clicked)
        layout.addWidget(tree_widget)

        close_btn = QPushButton("Закрыть")
        close_btn.clicked.connect(dialog.reject)
        layout.addWidget(close_btn)
        
        dialog.exec_()

    def show_save_dialog(self):
        if not self.selected_files:
            QMessageBox.warning(self, "Внимание", "Сначала загрузите файлы.")
            return

        dialog = QDialog(self)
        dialog.setWindowTitle("Сохранить")
        dialog.resize(500, 420)
        dialog.setStyleSheet(self._build_stylesheet())
        layout = QVBoxLayout(dialog)

        save_structure_cb = QCheckBox("Сохранить с сохранением структуры")
        save_map_cb = QCheckBox("Сохранить карту соответствий")
        save_structure_cb.setChecked(True)
        save_map_cb.setChecked(True)
        layout.addWidget(save_structure_cb)
        layout.addWidget(save_map_cb)

        tree_widget = QTreeWidget()
        tree_widget.setHeaderHidden(True)
        tree_widget.setColumnCount(1)
        tree_widget.setAlternatingRowColors(False)
        tree_widget.setUniformRowHeights(True)
        tree_widget.setRootIsDecorated(True)

        def build_file_tree():
            tree_widget.clear()
            if save_structure_cb.isChecked():
                root_item = QTreeWidgetItem(tree_widget, ["Файлы"])
                root_item.setExpanded(True)
                root_item.setFlags(root_item.flags() & ~Qt.ItemIsUserCheckable)
                folder_map = {}
                for file_path in self.selected_files:
                    if file_path == self.VIRTUAL_TEXT_PATH:
                        item = QTreeWidgetItem(root_item, ["my text"])
                        item.setCheckState(0, Qt.Checked)
                        item.setData(0, Qt.UserRole, file_path)
                        continue
                    
                    rel_name = self.file_relative_paths.get(file_path, os.path.basename(file_path)).replace('\\', '/')
                    parts = [p for p in rel_name.split('/') if p]
                    parent_item = root_item
                    for idx, part in enumerate(parts[:-1]):
                        current_key = '/'.join(parts[:idx + 1])
                        child = folder_map.get(current_key)
                        if child is None:
                            child = QTreeWidgetItem(parent_item, [part])
                            child.setExpanded(True)
                            child.setFlags(child.flags() & ~Qt.ItemIsUserCheckable)
                            folder_map[current_key] = child
                        parent_item = child

                    leaf_name = parts[-1] if parts else os.path.basename(file_path)
                    leaf_item = QTreeWidgetItem(parent_item, [leaf_name])
                    leaf_item.setCheckState(0, Qt.Checked)
                    leaf_item.setData(0, Qt.UserRole, file_path)
                    leaf_item.setData(0, Qt.UserRole + 1, rel_name)
            else:
                for file_path in self.selected_files:
                    if file_path == self.VIRTUAL_TEXT_PATH:
                        item = QTreeWidgetItem(tree_widget, ["my text"])
                    else:
                        rel_name = self.file_relative_paths.get(file_path, os.path.basename(file_path)).replace('\\', '/')
                        item = QTreeWidgetItem(tree_widget, [os.path.basename(rel_name)])
                    item.setCheckState(0, Qt.Checked)
                    item.setData(0, Qt.UserRole, file_path)

        def collect_selected_files():
            selected_paths = []
            root = tree_widget.invisibleRootItem()
            stack = [root]
            while stack:
                parent = stack.pop()
                for idx in range(parent.childCount()):
                    child = parent.child(idx)
                    if child.childCount() > 0:
                        stack.append(child)
                        continue
                    if child.checkState(0) == Qt.Checked:
                        file_path = child.data(0, Qt.UserRole)
                        if file_path:
                            selected_paths.append(file_path)
            return selected_paths

        build_file_tree()
        save_structure_cb.toggled.connect(lambda _: build_file_tree())
        layout.addWidget(tree_widget)

        buttons = QHBoxLayout()
        ok_btn = QPushButton("Сохранить")
        cancel_btn = QPushButton("Отмена")
        buttons.addStretch()
        buttons.addWidget(ok_btn)
        buttons.addWidget(cancel_btn)
        layout.addLayout(buttons)

        def save_selected():
            selected_paths = collect_selected_files()
            if not selected_paths:
                QMessageBox.warning(dialog, "Внимание", "Выберите хотя бы один файл для сохранения.")
                return

            folder_path = QFileDialog.getExistingDirectory(dialog, "Выберите папку для сохранения")
            if not folder_path:
                return

            if save_map_cb.isChecked():
                map_name = 'mapping.json'
                map_path = os.path.join(folder_path, map_name)
                with open(map_path, 'w', encoding='utf-8') as dst:
                    json.dump(self.current_mapping or {}, dst, ensure_ascii=False, indent=2)

            for file_path in selected_paths:
                content = self.anonymized_file_texts.get(file_path)
                if content is None:
                    content = self.file_text_cache.get(file_path) or self._extract_text_from_file(file_path)

                if file_path == self.VIRTUAL_TEXT_PATH:
                    rel_name = "my_text.txt"
                    safe_name = "my_text_anon.txt"
                else:
                    rel_name = self.file_relative_paths.get(file_path, os.path.basename(file_path)).replace('\\', '/')
                    file_name = os.path.basename(rel_name)
                    base_name, ext = os.path.splitext(file_name)
                    safe_name = f"{base_name}_anon.txt"

                if save_structure_cb.isChecked() and file_path != self.VIRTUAL_TEXT_PATH:
                    target_dir = os.path.join(folder_path, os.path.dirname(rel_name))
                    os.makedirs(target_dir, exist_ok=True)
                    target_path = os.path.join(target_dir, safe_name)
                else:
                    target_path = os.path.join(folder_path, safe_name)

                with open(target_path, 'w', encoding='utf-8') as dst:
                    dst.write(content or '')

            QMessageBox.information(dialog, "Готово", f"Сохранено файлов: {len(selected_paths)}")
            dialog.close()

        ok_btn.clicked.connect(save_selected)
        cancel_btn.clicked.connect(dialog.reject)
        dialog.exec_()

    def clear_selected_files(self):
        if self.active_file_path and not self.text_input_editor.isReadOnly():
            self.file_text_cache[self.active_file_path] = self.text_input_editor.toPlainText()

        self.file_list_widget.clear()
        self.selected_files.clear()
        self.current_file_replacements = {}
        self.current_mapping = {}
        self.last_anonymized_result = None
        self.current_map_path = None
        self.current_map_name = None
        self.deanon_reverse_map = {}
        self.file_text_cache.clear()
        self.anonymized_file_texts.clear()
        self.base_mapping_snapshot = {}
        self.excluded_aliases.clear()
        self.active_file_path = None
        self.text_input_editor.setReadOnly(False)
        self.text_input_editor.clear()
        self.text_input_editor.setPlainText("")
        self.text_input_editor.setHtml("")
        self._clear_editor_selection()
        self.mapping_table.blockSignals(True)
        self.mapping_table.clearSelection()
        self.mapping_table.setCurrentCell(-1, -1)
        self.mapping_table.blockSignals(False)
        self.mapping_table.clearContents()
        self.mapping_table.setRowCount(0)
        self.map_combo.blockSignals(True)
        self.map_combo.clear()
        self.map_combo.blockSignals(False)
        self.preview_browser.setHtml("<p style='color:#64748b;'>Список файлов очищен. Загрузите документ для предпросмотра.</p>")
        self.selected_mapping_alias = None
        self.file_relative_paths.clear()
        self.update_file_count()
        self.refresh_controls()

    def _clear_editor_selection(self):
        cursor = self.text_input_editor.textCursor()
        cursor.clearSelection()
        cursor.setPosition(0)
        self.text_input_editor.setTextCursor(cursor)

    def refresh_controls(self):
        has_files = bool(self.selected_files)
        current_text = self.text_input_editor.toPlainText().strip()

        if not has_files and not current_text:
            self.text_input_editor.setPlaceholderText("Вставьте или начните печатать")
        else:
            self.text_input_editor.setPlaceholderText("")

        self.btn_anonymize.setEnabled(has_files or bool(current_text))
        self.btn_structure.setEnabled(has_files)
        self.btn_save_results.setEnabled(has_files and (self.last_anonymized_result is not None or bool(self.anonymized_file_texts)))
        self.btn_deanonymize.setEnabled(True)
        self.btn_load.setEnabled(True)
        self.btn_clear.setEnabled(True)

    def _file_display_name(self, file_path):
        if not file_path:
            return ""
        if file_path == self.VIRTUAL_TEXT_PATH:
            return "my text"
        return os.path.basename(file_path) or file_path

    def _get_active_file_path(self):
        if self.active_file_path and os.path.exists(self.active_file_path):
            return self.active_file_path
        if self.selected_files:
            return self.selected_files[0]
        return None

    def _collect_files_from_drag_paths(self, paths):
        supported = {'.docx', '.pdf', '.jpg', '.jpeg', '.png'}
        collected = []
        for raw_path in paths:
            if not raw_path or not os.path.exists(raw_path):
                continue
            if os.path.isdir(raw_path):
                for root, _, files in os.walk(raw_path):
                    for file_name in files:
                        full_path = os.path.join(root, file_name)
                        if os.path.splitext(full_path)[1].lower() in supported:
                            relative_path = os.path.relpath(full_path, raw_path)
                            collected.append((full_path, relative_path))
            else:
                ext = os.path.splitext(raw_path)[1].lower()
                if ext in supported:
                    collected.append((raw_path, os.path.basename(raw_path)))

        unique = []
        seen = set()
        for full_path, rel_path in collected:
            if full_path not in seen:
                seen.add(full_path)
                unique.append((full_path, rel_path))
        return unique

    def _add_files_from_paths(self, file_paths):
        valid_files = []
        for file_path in file_paths:
            if not file_path or not os.path.exists(file_path):
                continue
            ext = os.path.splitext(file_path)[1].lower()
            if ext in {'.docx', '.pdf', '.jpg', '.jpeg', '.png'}:
                valid_files.append(file_path)

        if not valid_files:
            QMessageBox.warning(self, "Внимание", "Перетащите файлы поддерживаемых форматов: DOCX, PDF, PNG, JPG.")
            return

        relative_map = {}
        for file_path in valid_files:
            relative_map[file_path] = os.path.basename(file_path)

        for file_path in valid_files:
            self.add_file_to_list(file_path, relative_map.get(file_path))

        final_file = valid_files[-1]
        self.file_list_widget.setCurrentRow(self.file_list_widget.count() - 1)
        self.show_file_preview(final_file)
        self.populate_text_editor_from_file(final_file)
        self.refresh_controls()

    def _reset_drag_style(self):
        self.text_input_editor.setStyleSheet("")

    def eventFilter(self, obj, event):
        if obj is self.text_input_editor:
            if event.type() == QEvent.MouseButtonRelease and event.button() == Qt.LeftButton:
                anchor = self.text_input_editor.anchorAt(event.pos())
                if anchor.startswith('entity:'):
                    alias = anchor.split(':', 1)[1]
                    self.selected_mapping_alias = alias
                    self._sync_mapping_selection_to_active_alias()
                    self._scroll_to_first_occurrence(alias)
                    return True

            if event.type() in (QEvent.DragEnter, QEvent.Drop, QEvent.DragLeave):
                mime = event.mimeData()
                if mime and mime.hasUrls():
                    if event.type() == QEvent.DragEnter:
                        self.text_input_editor.setStyleSheet("QTextEdit { border: 2px dashed #4f46e5; background: #f8faff; }")
                        event.acceptProposedAction()
                        return True
                    if event.type() == QEvent.DragLeave:
                        self._reset_drag_style()
                        return True
                    paths = [url.toLocalFile() for url in mime.urls() if url.toLocalFile()]
                    if paths:
                        collected = self._collect_files_from_drag_paths(paths)
                        for full_path, rel_path in collected:
                            self.add_file_to_list(full_path, rel_path)
                        if collected:
                            last_path = collected[-1][0]
                            self.show_file_preview(last_path)
                            self.populate_text_editor_from_file(last_path)
                        self._reset_drag_style()
                        event.acceptProposedAction()
                        return True
        return super().eventFilter(obj, event)

    def dragEnterEvent(self, event):
        if event.mimeData().hasUrls():
            self.text_input_editor.setStyleSheet("QTextEdit { border: 2px dashed #4f46e5; background: #f8faff; }")
            event.acceptProposedAction()

    def dropEvent(self, event):
        if not event.mimeData().hasUrls():
            event.ignore()
            return

        paths = [url.toLocalFile() for url in event.mimeData().urls() if url.toLocalFile()]
        if paths:
            collected = self._collect_files_from_drag_paths(paths)
            for full_path, rel_path in collected:
                self.add_file_to_list(full_path, rel_path)
            if collected:
                last_path = collected[-1][0]
                self.show_file_preview(last_path)
                self.populate_text_editor_from_file(last_path)
            self._reset_drag_style()
            event.acceptProposedAction()

    def load_single_file(self):
        file_path, _ = QFileDialog.getOpenFileName(
            self, "Выберите файл", "",
            "Все файлы (*.docx *.pdf *.jpg *.png);;DOCX (*.docx);;PDF (*.pdf);;Изображения (*.jpg *.jpeg *.png)"
        )
        if file_path:
            self.add_file_to_list(file_path, os.path.basename(file_path))
            self.show_file_preview(file_path)
            self.populate_text_editor_from_file(file_path)
            self.refresh_controls()

    def load_folder(self):
        folder_path = QFileDialog.getExistingDirectory(self, "Выберите папку")
        if folder_path:
            for root, _, files in os.walk(folder_path):
                for file in files:
                    if file.endswith(('.docx', '.pdf', '.jpg', '.png')):
                        full_path = os.path.join(root, file)
                        relative_path = os.path.relpath(full_path, folder_path)
                        self.add_file_to_list(full_path, relative_path)
            if self.selected_files:
                last_path = self.selected_files[-1]
                self.show_file_preview(last_path)
                self.populate_text_editor_from_file(last_path)
            self.refresh_controls()

    def add_file_to_list(self, file_path, relative_path=None):
        if file_path not in self.selected_files:
            self.selected_files.append(file_path)
            if relative_path is None:
                relative_path = os.path.basename(file_path)
            self.file_relative_paths[file_path] = relative_path
            list_item = QListWidgetItem(self._file_display_name(file_path))
            list_item.setData(Qt.UserRole, file_path)
            self.file_list_widget.addItem(list_item)
            self.update_file_count()
            self.statusbar.showMessage(f"Файл загружен: {self._file_display_name(file_path)}")
            self.refresh_controls()

    def _clear_text_editor_if_manual_input(self):
        if self.text_input_editor.toPlainText().strip():
            return
        self.file_list_widget.clear()
        self.selected_files.clear()
        self.current_file_replacements = {}
        self.preview_file_label.setText("Файл: не выбран")
        self.preview_browser.setHtml("<p style='color:#64748b;'>Текст вставлен вручную. Для обработки используйте введённый текст.</p>")
        self.update_file_count()
        self.refresh_controls()

    def update_file_count(self):
        self.info_label.setText(f"Загружено файлов: {len(self.selected_files)}")
        self.file_tab_bar.blockSignals(True)
        while self.file_tab_bar.count() > 0:
            self.file_tab_bar.removeTab(self.file_tab_bar.count() - 1)
        for file_path in self.selected_files:
            self.file_tab_bar.addTab(self._file_display_name(file_path))
            self.file_tab_bar.setTabData(self.file_tab_bar.count() - 1, file_path)
        if self.selected_files:
            active_path = self.active_file_path if self.active_file_path in self.selected_files else self.selected_files[0]
            for idx in range(self.file_tab_bar.count()):
                if self.file_tab_bar.tabData(idx) == active_path:
                    self.file_tab_bar.setCurrentIndex(idx)
                    break
        self.file_tab_bar.blockSignals(False)

    def on_file_tab_changed(self, index):
        if index < 0 or self.file_tab_bar.count() == 0:
            return
        file_path = self.file_tab_bar.tabData(index)
        if not file_path or file_path not in self.selected_files:
            return
        self.active_file_path = file_path
        if self.file_list_widget.count() > self.selected_files.index(file_path):
            self.file_list_widget.setCurrentRow(self.selected_files.index(file_path))
        self.show_file_preview(file_path)
        self._render_current_active_file_in_editor()
        self.statusbar.showMessage(f"Выбран файл: {self._file_display_name(file_path)}")
        self.refresh_controls()

    def on_file_tab_close_requested(self, index):
        if index < 0 or index >= self.file_tab_bar.count():
            return
        file_path = self.file_tab_bar.tabData(index)
        if not file_path:
            return

        if file_path in self.selected_files:
            self.selected_files.remove(file_path)
        self.file_text_cache.pop(file_path, None)
        self.preview_text_cache.pop(file_path, None)
        self.anonymized_file_texts.pop(file_path, None)
        self.current_file_replacements.pop(file_path, None)
        self.file_relative_paths.pop(file_path, None)
        if self.active_file_path == file_path:
            self.active_file_path = self.selected_files[0] if self.selected_files else None

        if self.file_list_widget.count():
            self.file_list_widget.clear()
            for item_path in self.selected_files:
                list_item = QListWidgetItem(self._file_display_name(item_path))
                list_item.setData(Qt.UserRole, item_path)
                self.file_list_widget.addItem(list_item)

        self.update_file_count()

        if self.selected_files:
            self.show_file_preview(self.active_file_path)
            self._render_current_active_file_in_editor()
            if self.last_anonymized_result is not None and self.current_mapping:
                self.start_anonymization()
        else:
            self.active_file_path = None
            self.text_input_editor.setReadOnly(False)
            self.text_input_editor.clear()
            self.text_input_editor.setPlainText("")
            self.text_input_editor.setPlaceholderText("Вставьте или начните печатать")
            self.preview_browser.setHtml("<p style='color:#64748b;'>Нет загруженных файлов. Загрузите документ для просмотра.</p>")
            self.mapping_table.clearContents()
            self.mapping_table.setRowCount(0)
            self.current_mapping = {}
            self.base_mapping_snapshot = {}
            self.current_file_replacements = {}
            self.anonymized_file_texts.clear()
            self.last_anonymized_result = None

        self.refresh_controls()

    def refresh_map_list(self):
        try:
            maps = self.map_manager.list_maps()
        except Exception:
            maps = []

        self.available_maps = maps
        self.map_combo.blockSignals(True)
        self.map_combo.clear()
        for item in maps:
            label = item.get('name') or os.path.basename(item.get('path', ''))
            self.map_combo.addItem(label, item.get('path'))
        self.map_combo.blockSignals(False)

        if self.current_map_path:
            idx = self.map_combo.findData(self.current_map_path)
            if idx >= 0:
                self.map_combo.setCurrentIndex(idx)
        elif self.map_combo.count():
            self.map_combo.setCurrentIndex(0)

    def on_map_combo_changed(self, index):
        if index < 0 or self.map_combo.count() == 0:
            return
        path = self.map_combo.itemData(index)
        if not path:
            return
        self.current_map_path = path
        try:
            mapping = self.map_manager.load_map(path)
            self.activate_mapping(mapping, source_path=path)
        except Exception:
            pass

    def activate_mapping(self, mapping, source_path=None, source_name=None):
        if not mapping:
            return
        
        user_entities = {alias: info for alias, info in self.current_mapping.items()
                        if info.get('category') == 'USER'}
        mapping.update(user_entities)
        
        self.current_mapping = mapping
        self.current_map_path = source_path or self.current_map_path
        self.current_map_name = source_name or os.path.basename(self.current_map_path) if self.current_map_path else None
        self.deanon_reverse_map = Deanon().load_map_from_data(mapping)
        self.render_mapping_table_from_current_mapping()
        self.statusbar.showMessage(f"Активна карта: {self.current_map_name or 'выбранная карта'}")

    def _normalize_file_list(self, file_paths):
        return ", ".join(self._file_display_name(path) for path in (file_paths or []))

    def _get_active_mapping(self):
        base = self.base_mapping_snapshot or self.current_mapping
        if not base:
            return {}
        return {alias: info for alias, info in base.items() if alias not in self.excluded_aliases}

    def on_mapping_exclusion_toggled(self, alias, checked):
        if checked:
            self.excluded_aliases.discard(alias)
        else:
            self.excluded_aliases.add(alias)

        self.render_mapping_table_from_current_mapping()

        if not self.selected_files and not self.text_input_editor.toPlainText().strip():
            return

        self._suppress_processing_notifications = True
        try:
            if self.selected_files:
                self.start_anonymization()
            elif self.text_input_editor.toPlainText().strip():
                self.anonymize_text_content(self.text_input_editor.toPlainText())
        finally:
            self._suppress_processing_notifications = False

    def render_mapping_table_from_current_mapping(self):
        mapping_source = self.base_mapping_snapshot or self.current_mapping or {}
        user_entities = {alias: info for alias, info in self.current_mapping.items()
                        if info.get('category') == 'USER'}
        merged_mapping = dict(mapping_source)
        merged_mapping.update(user_entities)
        
        mapping_list = []
        for alias, info in merged_mapping.items():
            mapping_list.append({
                "pseudonym": alias,
                "original": info.get("original", "") if isinstance(info, dict) else str(info),
                "category": info.get("category", "") if isinstance(info, dict) else "",
                "file": ", ".join(self._file_display_name(fp) for fp in (info.get("files", []) if isinstance(info, dict) else []))
            })

        self.mapping_table.clearContents()
        self.mapping_table.setRowCount(len(mapping_list))
        for i, item in enumerate(mapping_list):
            row_alias = item["pseudonym"]
            is_excluded = row_alias in self.excluded_aliases

            self.mapping_table.setItem(i, 0, QTableWidgetItem(item["pseudonym"]))
            self.mapping_table.setItem(i, 1, QTableWidgetItem(item["original"]))
            self.mapping_table.setItem(i, 2, QTableWidgetItem(item["category"]))
            self.mapping_table.setItem(i, 3, QTableWidgetItem(item["file"]))

            for col in range(4):
                cell = self.mapping_table.item(i, col)
                if cell:
                    cell.setBackground(Qt.white if not is_excluded else QColor(230, 230, 230))
                    cell.setForeground(Qt.black if not is_excluded else QColor(120, 120, 120))

            checkbox = QCheckBox("Вкл")
            checkbox.setChecked(not is_excluded)
            checkbox.toggled.connect(lambda checked, alias=row_alias: self.on_mapping_exclusion_toggled(alias, checked))
            self.mapping_table.setCellWidget(i, 4, checkbox)

    def load_map_from_file(self):
        file_path, _ = QFileDialog.getOpenFileName(
            self,
            "Выберите JSON с картой соответствия",
            "",
            "JSON файл (*.json)"
        )
        if not file_path:
            return
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                mapping = json.load(f)
            deanon = Deanon()
            normalized = deanon._normalize_mapping(mapping)
            self.activate_mapping(normalized, source_path=file_path, source_name=os.path.basename(file_path))
            self.refresh_map_list()
            self.statusbar.showMessage(f"Загружена карта: {os.path.basename(file_path)}")
            QMessageBox.information(self, "Карта загружена", f"Активна карта: {os.path.basename(file_path)}")
        except Exception as exc:
            QMessageBox.critical(self, "Ошибка карты", f"Не удалось загрузить карту: {exc}")

    def save_anonymized_documents(self):
        if not self.selected_files:
            QMessageBox.warning(self, "Внимание", "Сначала загрузите документ.")
            return

        if self.last_anonymized_result is None and not self.anonymized_file_texts:
            QMessageBox.warning(self, "Внимание", "Сначала анонимизируйте документ, а потом сохраняйте результат.")
            return

        if self.save_all_checkbox.isChecked():
            folder_path = QFileDialog.getExistingDirectory(
                self,
                "Выберите папку для сохранения всех анонимизированных документов"
            )
            if not folder_path:
                return

            saved = []
            for file_path in self.selected_files:
                content = self.anonymized_file_texts.get(file_path)
                if content is None:
                    continue
                base_name = os.path.splitext(self._file_display_name(file_path))[0]
                save_path = os.path.join(folder_path, f"{base_name}_anon.txt")
                with open(save_path, 'w', encoding='utf-8') as dst:
                    dst.write(content)
                saved.append(save_path)

            if saved:
                self.statusbar.showMessage(f"Сохранено документов: {len(saved)}")
                QMessageBox.information(self, "Готово", f"Сохранено файлов: {len(saved)}\nПапка: {folder_path}")
            else:
                QMessageBox.warning(self, "Внимание", "Нет данных для сохранения в выбранной папке.")
            return

        current_file = self._get_active_file_path() or self.selected_files[0]
        default_name = os.path.splitext(os.path.basename(current_file))[0] + "_anon.txt"
        default_path = os.path.join(os.path.dirname(current_file), default_name)

        save_path, _ = QFileDialog.getSaveFileName(
            self,
            "Сохранить анонимизированный документ",
            default_path,
            "Text files (*.txt)"
        )
        if not save_path:
            return

        try:
            if self.anonymized_file_texts and current_file in self.anonymized_file_texts:
                content = self.anonymized_file_texts[current_file]
            elif self.text_input_editor.isReadOnly():
                content = self.text_input_editor.toPlainText()
            else:
                content = self.text_input_editor.toPlainText().strip()

            if not content:
                content = self.last_anonymized_result.get('last_text') or ''

            with open(save_path, 'w', encoding='utf-8') as dst:
                dst.write(content)

            self.statusbar.showMessage(f"Документ сохранён: {os.path.basename(save_path)}")
            QMessageBox.information(self, "Готово", f"Документ сохранён в:\n{save_path}")
        except Exception as exc:
            QMessageBox.critical(self, "Ошибка сохранения", str(exc))

    def select_save_folder_from_main(self):
        folder_path = QFileDialog.getExistingDirectory(
            self, "Выберите папку для сохранений результатов"
        )
        if folder_path:
            self.save_folder = folder_path
            self.refresh_controls()
            self.statusbar.showMessage(f"Папка сохранения установлена: {folder_path}")
            if self.selected_files:
                self.show_file_preview(self.selected_files[0])

    def _extract_text_from_file(self, file_path):
        if file_path == self.VIRTUAL_TEXT_PATH:
            return self.file_text_cache.get(file_path, "")
        
        if file_path in self.preview_text_cache:
            return self.preview_text_cache[file_path]

        ext = os.path.splitext(file_path)[1].lower()
        parser_map = {
            '.docx': DOCXParser(),
            '.pdf': PDFParser(),
            '.jpg': ImageParser(),
            '.jpeg': ImageParser(),
            '.png': ImageParser()
        }
        parser = parser_map.get(ext)
        if parser is None:
            return ""

        try:
            doc = parser.parse(file_path)
        except Exception:
            return ""

        chunks = []
        for page in doc.pages:
            for para in page.paragraphs:
                text = ''.join(f.text for f in para.fragments)
                if text.strip():
                    chunks.append(text.strip())
            for table in page.tables:
                for row in table.rows:
                    for cell in row:
                        for para in cell.paragraphs:
                            text = ''.join(f.text for f in para.fragments)
                            if text.strip():
                                chunks.append(text.strip())

        result = "\n".join(chunks)
        self.preview_text_cache[file_path] = result
        return result

    def _make_alias(self, category, index=1):
        mapping = {
            'PHONE': 'Phone',
            'EMAIL': 'Email',
            'PASSPORT': 'Passport',
            'INN': 'INN',
            'BANK_ACCOUNT': 'BankAccount',
            'BIC': 'BIC',
            'DEPT_CODE': 'DeptCode',
            'PERSON': 'Person',
            'ORG': 'Org',
            'ADDRESS': 'Address'
        }
        base = mapping.get(category, category.capitalize())
        suffix = '' if index <= 1 else str(index)
        return f"[{base}{suffix}]"

    def _highlight_replacements_in_html(self, text, replacements, selected_value=None):
        text = text.replace('\r\n', '\n').replace('\r', '\n').replace('\t', '    ')

        def htmlize(value):
            escaped = html.escape(value).replace(' ', '&nbsp;').replace('\n', '<br>')
            return f"<span style='color:#111827;'>{escaped}</span>"

        def find_all_occurrences(block_text, needle):
            needle = str(needle).strip()
            if not needle:
                return []
            lowered = block_text.lower()
            query = needle.lower()
            results = []
            start = 0
            while True:
                idx = lowered.find(query, start)
                if idx == -1:
                    break
                results.append((idx, idx + len(needle)))
                start = idx + 1
            return results

        def render_block(block_text, block_replacements, block_selected_value=None):
            if not block_replacements and not block_selected_value:
                return htmlize(block_text)

            local_selected = []
            if block_selected_value:
                needle = str(block_selected_value).strip()
                local_selected.extend(find_all_occurrences(block_text, needle))

            dynamic_replacements = []
            for rep in block_replacements:
                value = str(rep.get('value') or rep.get('original') or '').strip()
                if not value:
                    start = int(rep.get('start', 0))
                    end = int(rep.get('end', start))
                    if 0 <= start < len(block_text) and 0 <= end <= len(block_text):
                        dynamic_replacements.append((start, end, 'replacement', rep.get('alias', '')))
                    continue

                for start, end in find_all_occurrences(block_text, value):
                    if any(start < s_end and end > s_start for s_start, s_end in local_selected):
                        continue
                    dynamic_replacements.append((start, end, 'replacement', rep.get('alias', '')))

            ranges = []
            for start, end, kind, label in dynamic_replacements:
                ranges.append((start, end, kind, label))
            for start, end in local_selected:
                ranges.append((start, end, 'selected', block_selected_value))

            if not ranges:
                return htmlize(block_text)

            ranges.sort(key=lambda item: (item[0], item[2] == 'selected'))

            segments = []
            cursor = 0
            for start, end, kind, label in ranges:
                if start < cursor:
                    continue
                if start > len(block_text):
                    continue
                segments.append(htmlize(block_text[cursor:start]))
                fragment = htmlize(block_text[start:end])
                if kind == 'selected':
                    segments.append(f"<span style='background:#dbeafe; color:#1e3a8a; border:1px solid #93c5fd; border-radius:6px; padding:0 3px; font-weight:700;'>{fragment}</span>")
                else:
                    alias = html.escape(str(label))
                    segments.append(f"<span style='background:#fee2e2; color:#991b1b; border-radius:6px; padding:0 3px; font-weight:700;'>{alias}</span>")
                cursor = end
            if cursor < len(block_text):
                segments.append(htmlize(block_text[cursor:]))
            return ''.join(segments)

        blocks = text.split('\n')
        if not blocks:
            blocks = [text]

        rendered = []
        for block in blocks:
            block_replacements = []
            for rep in replacements or []:
                value = str(rep.get('value') or rep.get('original') or '').strip()
                if not value:
                    continue
                if value.lower() in block.lower():
                    block_replacements.append(rep)
            rendered.append(
                f"<div style='white-space: pre-wrap; line-height: 1.8; font-family: Consolas, monospace; margin-bottom: 4px;'>"
                f"{render_block(block, block_replacements, selected_value)}"
                f"</div>"
            )

        if not rendered:
            return "<div style='white-space: pre-wrap; line-height: 1.8; font-family: Consolas, monospace;'> </div>"

        return ''.join(rendered)

    def _highlight_aliases_in_html(self, text, alias_map, selected_alias=None):
        text = text.replace('\r\n', '\n').replace('\r', '\n').replace('\t', '    ')
        aliases = sorted(alias_map.keys(), key=len, reverse=True)
        if not aliases:
            escaped = html.escape(text).replace('\n', '<br>')
            return f"<span style='color:#111827;'>{escaped}</span>"

        pattern = re.compile('|'.join(re.escape(alias) for alias in aliases))
        parts = []
        cursor = 0
        for match in pattern.finditer(text):
            start, end = match.span()
            normal_text = html.escape(text[cursor:start]).replace('\n', '<br>')
            if normal_text:
                parts.append(f"<span style='color:#111827;'>{normal_text}</span>")
            alias = match.group(0)
            if selected_alias and alias == selected_alias:
                parts.append(f"<a href='entity:{html.escape(alias)}' style='text-decoration:none; color:inherit;'><span style='background:#dbeafe; color:#1e3a8a; border:1px solid #93c5fd; border-radius:6px; padding:0 3px; font-weight:700;'>{html.escape(alias)}</span></a>")
            else:
                parts.append(f"<a href='entity:{html.escape(alias)}' style='text-decoration:none; color:inherit;'><span style='background:#fef3c7; color:#78350f; border-radius:6px; padding:0 3px; font-weight:700;'>{html.escape(alias)}</span></a>")
            cursor = end
        final_text = html.escape(text[cursor:]).replace('\n', '<br>')
        if final_text:
            parts.append(f"<span style='color:#111827;'>{final_text}</span>")
        return ''.join(parts)

    def on_editor_text_changed(self):
        if self.text_input_editor.isReadOnly():
            return
        
        current_text = self.text_input_editor.toPlainText()
        
        if self.active_file_path:
            self.file_text_cache[self.active_file_path] = current_text
        elif current_text.strip():
            if self.VIRTUAL_TEXT_PATH not in self.selected_files:
                self.selected_files.insert(0, self.VIRTUAL_TEXT_PATH)
                self.file_relative_paths[self.VIRTUAL_TEXT_PATH] = "my text"
                self.active_file_path = self.VIRTUAL_TEXT_PATH
                self.update_file_count()
            self.file_text_cache[self.VIRTUAL_TEXT_PATH] = current_text

    def on_editor_context_menu(self, position):
        menu = QMenu()
        menu.addAction("Копировать", lambda: self.text_input_editor.copy())
        menu.addAction("Выделить всё", lambda: self.text_input_editor.selectAll())
        
        cursor = self.text_input_editor.textCursor()
        if cursor.hasSelection():
            menu.addSeparator()
            menu.addAction("Сделать анонимом", self.on_make_anonymous)

        menu.exec_(self.text_input_editor.mapToGlobal(position))

    def on_make_anonymous(self):
        cursor = self.text_input_editor.textCursor()
        if not cursor.hasSelection():
            QMessageBox.warning(self, "Внимание", "Выберите текст для анонимизации.")
            return

        selected_text = cursor.selectedText()
        if not selected_text.strip():
            QMessageBox.warning(self, "Внимание", "Выделенный текст пуст.")
            return

        dialog = QDialog(self)
        dialog.setWindowTitle("Сделать анонимом")
        dialog.setModal(True)
        dialog.resize(400, 150)

        layout = QVBoxLayout()
        label = QLabel("Введите маску для замены (например, [ПОЧТА_1]):")
        input_field = QLineEdit()
        input_field.setPlaceholderText("[ПОЧТА_1]")

        button_layout = QHBoxLayout()
        ok_button = QPushButton("OK")
        cancel_button = QPushButton("Отмена")
        button_layout.addWidget(ok_button)
        button_layout.addWidget(cancel_button)

        layout.addWidget(label)
        layout.addWidget(input_field)
        layout.addLayout(button_layout)
        dialog.setLayout(layout)

        def apply_mask():
            mask = input_field.text().strip()
            if not mask:
                QMessageBox.warning(dialog, "Внимание", "Маска не может быть пустой.")
                return

            if not mask.startswith('['):
                mask = '[' + mask
            if not mask.endswith(']'):
                mask = mask + ']'

            if self.current_mapping and mask in self.current_mapping:
                QMessageBox.warning(dialog, "Внимание", f"Маска '{mask}' уже существует в карте соответствия.")
                return

            start = cursor.selectionStart()
            end = cursor.selectionEnd()
            current_text = self.text_input_editor.toPlainText()
            new_text = current_text[:start] + mask + current_text[end:]
            self.text_input_editor.setPlainText(new_text)

            if not self.current_mapping:
                self.current_mapping = {}

            file_path = self.active_file_path or self.VIRTUAL_TEXT_PATH
            self.current_mapping[mask] = {
                "original": selected_text,
                "category": "USER",
                "files": [file_path] if file_path else []
            }
            
            if not self.base_mapping_snapshot:
                self.base_mapping_snapshot = {}
            self.base_mapping_snapshot[mask] = self.current_mapping[mask]

            self.excluded_aliases.discard(mask)
            self.deanon_reverse_map = Deanon().load_map_from_data(self.current_mapping)
            
            self.mapping_table.blockSignals(True)
            self.render_mapping_table_from_current_mapping()
            self.mapping_table.blockSignals(False)

            new_cursor = self.text_input_editor.textCursor()
            new_cursor.setPosition(start)
            new_cursor.setPosition(start + len(mask), new_cursor.KeepAnchor)
            self.text_input_editor.setTextCursor(new_cursor)

            if self.current_mapping and self.text_input_editor.toPlainText().strip():
                self.text_input_editor.setHtml(
                    self._highlight_aliases_in_html(
                        self.text_input_editor.toPlainText(),
                        self.current_mapping,
                        selected_alias=mask
                    )
                )

            dialog.close()

        ok_button.clicked.connect(apply_mask)
        cancel_button.clicked.connect(dialog.close)
        input_field.returnPressed.connect(apply_mask)

        dialog.exec_()

    def _sync_mapping_selection_to_active_alias(self):
        if not self.selected_mapping_alias:
            self.mapping_table.blockSignals(True)
            self.mapping_table.clearSelection()
            self.mapping_table.setCurrentCell(-1, -1)
            self.mapping_table.blockSignals(False)
            return

        self.mapping_table.blockSignals(True)
        row_index = -1
        for row in range(self.mapping_table.rowCount()):
            alias_item = self.mapping_table.item(row, 0)
            if alias_item and alias_item.text() == self.selected_mapping_alias:
                row_index = row
                break
        if row_index >= 0:
            self.mapping_table.setCurrentCell(row_index, 0)
            self.mapping_table.selectRow(row_index)
        else:
            self.mapping_table.clearSelection()
            self.mapping_table.setCurrentCell(-1, -1)
        self.mapping_table.blockSignals(False)

    def _scroll_to_first_occurrence(self, alias):
        if not alias:
            return
        text = self.text_input_editor.toPlainText()
        if not text:
            return
        index = text.find(alias)
        if index < 0:
            return

        cursor = self.text_input_editor.textCursor()
        cursor.setPosition(index)
        self.text_input_editor.setTextCursor(cursor)
        self.text_input_editor.ensureCursorVisible()

    def on_mapping_header_clicked(self, logical_index):
        if logical_index != 4:
            return

        checkboxes = []
        for row in range(self.mapping_table.rowCount()):
            widget = self.mapping_table.cellWidget(row, 4)
            if isinstance(widget, QCheckBox):
                checkboxes.append(widget)

        if not checkboxes:
            return

        target_state = not all(cb.isChecked() for cb in checkboxes)
        for cb in checkboxes:
            cb.setChecked(target_state)

    def on_mapping_selection_changed(self):
        row = self.mapping_table.currentRow()
        current_editor_text = self.text_input_editor.toPlainText()

        if row < 0:
            self.selected_mapping_alias = None
            if self.current_mapping and current_editor_text.strip():
                self.text_input_editor.setHtml(self._highlight_aliases_in_html(current_editor_text, self.current_mapping, selected_alias=None))
            if self.selected_files:
                file_path = self.active_file_path or self.selected_files[0]
                preview_text = self._extract_text_from_file(file_path)
                if preview_text:
                    replacements = self._get_preview_replacements_for_file(file_path)
                    self.preview_browser.setHtml(self._highlight_replacements_in_html(preview_text, replacements, selected_value=None))
            return

        if not self.selected_files:
            self.selected_mapping_alias = None
            if self.current_mapping and current_editor_text.strip():
                self.text_input_editor.setHtml(self._highlight_aliases_in_html(current_editor_text, self.current_mapping, selected_alias=None))
            return

        original = self.mapping_table.item(row, 1).text() if self.mapping_table.item(row, 1) else ""
        selected_alias = None
        for alias, info in (self.current_mapping or {}).items():
            if str(info.get('original', '')).strip().lower() == str(original).strip().lower():
                selected_alias = alias
                break
        self.selected_mapping_alias = selected_alias

        if self.current_mapping and current_editor_text.strip():
            self.text_input_editor.setHtml(self._highlight_aliases_in_html(current_editor_text, self.current_mapping, selected_alias=selected_alias))

        if selected_alias:
            self._scroll_to_first_occurrence(selected_alias)

        file_path = self.active_file_path or self.selected_files[0]
        preview_text = self._extract_text_from_file(file_path)
        if preview_text:
            replacements = self._get_preview_replacements_for_file(file_path)
            self.preview_browser.setHtml(self._highlight_replacements_in_html(preview_text, replacements, selected_value=original))

    def _render_current_active_file_in_editor(self):
        if not self.active_file_path:
            return

        if self.current_mapping and self.active_file_path in self.anonymized_file_texts:
            self.text_input_editor.setReadOnly(True)
            self.text_input_editor.setHtml(self._highlight_aliases_in_html(
                self.anonymized_file_texts[self.active_file_path],
                self.current_mapping,
                selected_alias=self.selected_mapping_alias
            ))
            return

        self.text_input_editor.setReadOnly(False)
        text = self.file_text_cache.get(self.active_file_path)
        if text is None:
            text = self._extract_text_from_file(self.active_file_path)
            self.file_text_cache[self.active_file_path] = text
        self.text_input_editor.setPlainText(text or "")
        self._clear_editor_selection()
        if self.selected_mapping_alias and self.current_mapping and self.text_input_editor.toPlainText().strip():
            self.text_input_editor.setHtml(self._highlight_aliases_in_html(
                self.text_input_editor.toPlainText(),
                self.current_mapping,
                selected_alias=self.selected_mapping_alias
            ))

    def on_file_selected(self, item):
        if item is None:
            return

        if self.active_file_path and not self.text_input_editor.isReadOnly():
            self.file_text_cache[self.active_file_path] = self.text_input_editor.toPlainText()

        file_path = item.data(Qt.UserRole) or item.text()
        self.active_file_path = file_path
        self.selected_mapping_alias = None
        self.mapping_table.blockSignals(True)
        self.mapping_table.clearSelection()
        self.mapping_table.setCurrentCell(-1, -1)
        self.mapping_table.blockSignals(False)
        
        self.update_file_count()
        if os.path.exists(file_path):
            self.show_file_preview(file_path)
            self._render_current_active_file_in_editor()
            self._sync_mapping_selection_to_active_alias()
            self.statusbar.showMessage(f"Текст загруженного файла помещён в редактор: {self._file_display_name(file_path)}")
            self.refresh_controls()

    def populate_text_editor_from_file(self, file_path):
        self.active_file_path = file_path
        self.selected_mapping_alias = None
        self.mapping_table.blockSignals(True)
        self.mapping_table.clearSelection()
        self.mapping_table.setCurrentCell(-1, -1)
        self.mapping_table.blockSignals(False)
        
        text = self.file_text_cache.get(file_path)
        if text is None:
            text = self._extract_text_from_file(file_path)
            self.file_text_cache[file_path] = text
        self.text_input_editor.setReadOnly(False)
        self.text_input_editor.clear()
        self.text_input_editor.setPlainText(text or "")
        self._clear_editor_selection()
        self.update_file_count()
        self._sync_mapping_selection_to_active_alias()
        self.statusbar.showMessage(f"Текст загруженного файла помещён в редактор: {self._file_display_name(file_path)}")
        self.refresh_controls()

    def _get_preview_replacements_for_file(self, file_path):
        if hasattr(self, 'current_file_replacements'):
            return self.current_file_replacements.get(file_path, [])
        return []

    def show_file_preview(self, file_path, selected_value=None):
        if not file_path:
            return

        text = self._extract_text_from_file(file_path)
        if not text:
            self.preview_browser.setHtml("<p style='color:#64748b;'>Не удалось получить текст для предпросмотра этого файла.</p>")
            return

        if selected_value is None and self.selected_mapping_alias and self.current_mapping:
            selected_value = self.current_mapping.get(self.selected_mapping_alias, {}).get('original')

        replacement_candidates = self._get_preview_replacements_for_file(file_path)
        html_preview = self._highlight_replacements_in_html(text, replacement_candidates, selected_value=selected_value)
        self.preview_browser.setHtml(html_preview)

    def _build_global_anonymization_state(self, text_by_file):
        detector_manager = DetectorManager(self.settings)
        collected = []
        for file_path, text in text_by_file.items():
            for item in detector_manager.detect(text):
                value = str(item.get('value', '')).strip()
                if not value:
                    continue
                collected.append({
                    'file': file_path,
                    'value': value,
                    'category': item.get('category', 'UNKNOWN'),
                    'start': int(item.get('start', 0)),
                    'end': int(item.get('end', len(value)))
                })

        groups = {}
        counters = {}
        all_mapping = {}
        file_replacements = {}
        for item in collected:
            norm_key = (item['value'].lower().strip(), item['category'])
            groups.setdefault(norm_key, []).append(item)

        for (_, category), entries in groups.items():
            counters[category] = counters.get(category, 0) + 1
            alias = self._make_alias(category, counters[category])
            all_mapping[alias] = {
                'original': entries[0]['value'],
                'category': category,
                'files': sorted({entry['file'] for entry in entries})
            }
            for entry in entries:
                file_replacements.setdefault(entry['file'], []).append({
                    'start': entry['start'],
                    'end': entry['end'],
                    'value': entry['value'],
                    'category': category,
                    'alias': alias
                })

        filtered_mapping = {alias: info for alias, info in all_mapping.items() if alias not in self.excluded_aliases}
        filtered_file_replacements = {
            file_path: [rep for rep in reps if rep.get('alias') not in self.excluded_aliases]
            for file_path, reps in file_replacements.items()
        }

        anonymized_texts = {}
        for file_path, text in text_by_file.items():
            current = text
            for rep in sorted(filtered_file_replacements.get(file_path, []), key=lambda x: x['start'], reverse=True):
                current = current[:int(rep['start'])] + rep['alias'] + current[int(rep['end']):]
            anonymized_texts[file_path] = current

        return all_mapping, filtered_mapping, filtered_file_replacements, anonymized_texts

    def _add_user_replacements(self, file_replacements, texts_by_file, user_entities):
        """
        Добавляет replacements для пользовательских сущностей (категория 'USER')
        в словарь file_replacements на основе исходных текстов.
        """
        for alias, info in user_entities.items():
            if alias in self.excluded_aliases:
                continue
            original = info.get('original', '')
            if not original:
                continue
            files = info.get('files', [])
            for file_path in files:
                if file_path not in texts_by_file:
                    continue
                text = texts_by_file[file_path]
                if not text:
                    continue
                # Регистронезависимый поиск всех вхождений
                lowered_text = text.lower()
                lowered_original = original.lower()
                start = 0
                while True:
                    idx = lowered_text.find(lowered_original, start)
                    if idx == -1:
                        break
                    file_replacements.setdefault(file_path, []).append({
                        'start': idx,
                        'end': idx + len(original),
                        'value': original,
                        'category': 'USER',
                        'alias': alias
                    })
                    start = idx + 1
        return file_replacements

    def start_anonymization(self):
        editor_text = self.text_input_editor.toPlainText().strip()
        if editor_text and not self.selected_files:
            self.anonymize_text_content(editor_text)
            return

        if not self.selected_files:
            QMessageBox.warning(self, "Внимание", "Выберите файл или вставьте текст для обработки перед запуском!")
            return

        self.statusbar.showMessage("Начало обработки файлов...")
        self.btn_anonymize.setEnabled(False)
        self.btn_save_results.setEnabled(False)

        texts_by_file = {}
        for file_path in self.selected_files:
            if file_path == self.active_file_path and self.text_input_editor.isReadOnly() is False:
                texts_by_file[file_path] = self.text_input_editor.toPlainText()
            else:
                texts_by_file[file_path] = self.file_text_cache.get(file_path, self._extract_text_from_file(file_path))

        full_mapping, mapping, file_replacements, anonymized_texts = self._build_global_anonymization_state(texts_by_file)
        
        user_entities = {alias: info for alias, info in self.current_mapping.items()
                        if info.get('category') == 'USER'}
        full_mapping.update(user_entities)
        mapping.update(user_entities)
        file_replacements = self._add_user_replacements(file_replacements, texts_by_file, user_entities)

        # Пересчёт anonymized_texts с учётом новых замен
        anonymized_texts = {}
        for file_path, text in texts_by_file.items():
            current = text
            for rep in sorted(file_replacements.get(file_path, []), key=lambda x: x['start'], reverse=True):
                current = current[:rep['start']] + rep['alias'] + current[rep['end']:]
            anonymized_texts[file_path] = current

        self.base_mapping_snapshot = full_mapping
        self.current_mapping = mapping
        self.current_file_replacements = file_replacements
        self.anonymized_file_texts = anonymized_texts
        self.last_anonymized_result = {
            'mapping': mapping,
            'file_replacements': file_replacements,
            'files_processed': len(self.selected_files),
            'save_folder': os.getcwd(),
            'preserve_structure': self.preserve_structure,
            'last_text': anonymized_texts.get(self.active_file_path or self.selected_files[0], '')
        }

        if mapping:
            try:
                timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
                self.current_map_path = self.map_manager.save_map(mapping, name=f"map_{timestamp}")
            except Exception:
                self.current_map_path = None

        self.render_mapping_table_from_current_mapping()
        current_file = self.active_file_path or self.selected_files[0]
        if self.active_file_path is None:
            self.active_file_path = current_file
        self.show_file_preview(current_file)
        self._render_current_active_file_in_editor()
        self.statusbar.showMessage("Анонимизация завершена для всех выбранных файлов")
        if not self._suppress_processing_notifications:
            QMessageBox.information(self, "Готово", f"Обработано файлов: {len(self.selected_files)}\nНайдено сущностей: {len(mapping)}")

    def anonymize_text_content(self, text: str):
        if not text.strip():
            QMessageBox.warning(self, "Внимание", "Сначала введите или загрузите текст для анонимизации.")
            return

        detector_manager = DetectorManager(self.settings)
        entities = []
        for item in detector_manager.detect(text):
            value = str(item.get('value', '')).strip()
            if not value:
                continue
            entities.append({
                'value': value,
                'category': item.get('category', 'UNKNOWN'),
                'start': int(item.get('start', 0)),
                'end': int(item.get('end', len(value)))
            })

        if not entities:
            QMessageBox.information(self, "Результат", "Сущности не найдены. Текст оставлен без изменений.")
            self.text_input_editor.setReadOnly(False)
            self.text_input_editor.setPlainText(text)
            self.statusbar.showMessage("Сущности не найдены")
            return

        grouped = {}
        counters = {}
        for item in entities:
            key = (item['value'].lower(), item['category'])
            grouped.setdefault(key, []).append(item)

        mapping = {}
        replacements = []
        for (value, category), matches in grouped.items():
            counters[category] = counters.get(category, 0) + 1
            alias = self._make_alias(category, counters[category])
            mapping[alias] = {
                'original': matches[0]['value'],
                'category': category,
                'files': ['edited_text']
            }
            for match in matches:
                replacements.append({
                    'start': match['start'],
                    'end': match['end'],
                    'value': match['value'],
                    'category': category,
                    'alias': alias
                })

        anonymized_text = text
        for item in sorted(replacements, key=lambda x: x['start'], reverse=True):
            anonymized_text = anonymized_text[:item['start']] + item['alias'] + anonymized_text[item['end']:] 

        full_mapping = mapping.copy()
        filtered_replacements = []
        filtered_mapping = {}
        for item in replacements:
            alias = item.get('alias')
            if alias in self.excluded_aliases:
                continue
            filtered_replacements.append(item)
        for alias, info in full_mapping.items():
            if alias in self.excluded_aliases:
                continue
            filtered_mapping[alias] = info

        user_entities = {alias: info for alias, info in self.current_mapping.items()
                        if info.get('category') == 'USER'}
        full_mapping.update(user_entities)
        filtered_mapping.update(user_entities)

        for alias, info in user_entities.items():
            if alias in self.excluded_aliases:
                continue
            original = info.get('original', '')
            if not original:
                continue
            # Поиск всех вхождений (регистронезависимо)
            lowered_text = text.lower()
            lowered_original = original.lower()
            start = 0
            while True:
                idx = lowered_text.find(lowered_original, start)
                if idx == -1:
                    break
                filtered_replacements.append({
                    'start': idx,
                    'end': idx + len(original),
                    'value': original,
                    'category': 'USER',
                    'alias': alias
                })
                start = idx + 1

        self.base_mapping_snapshot = full_mapping
        self.current_mapping = filtered_mapping
        self.current_file_replacements = {'edited_text': filtered_replacements}
        self.deanon_reverse_map = Deanon().load_map_from_data(filtered_mapping)

        rebuilt_text = text
        for rep in sorted(filtered_replacements, key=lambda x: x['start'], reverse=True):
            rebuilt_text = rebuilt_text[:int(rep['start'])] + rep['alias'] + rebuilt_text[int(rep['end']):]

        self.last_anonymized_result = {
            'mapping': filtered_mapping,
            'file_replacements': {'edited_text': filtered_replacements},
            'files_processed': 1,
            'save_folder': os.getcwd(),
            'preserve_structure': self.preserve_structure,
            'last_text': rebuilt_text
        }

        highlighted_html = self._highlight_aliases_in_html(rebuilt_text, filtered_mapping)
        self.text_input_editor.setReadOnly(True)
        self.text_input_editor.setHtml(highlighted_html)
        self.render_mapping_table_from_current_mapping()
        self.statusbar.showMessage("Текст анонимизирован")
        if not self._suppress_processing_notifications:
            QMessageBox.information(self, "Готово", f"Текст анонимизирован. Найдено сущностей: {len(mapping)}")

    def update_progress(self, value, text, file_index):
        self.statusbar.showMessage(text)

    def on_anonymization_finished(self, result_data):
        self.statusbar.showMessage("Обработка завершена")
        self.btn_anonymize.setEnabled(bool(self.selected_files))
        self.btn_save_results.setEnabled(bool(self.selected_files))

        if "error" in result_data:
            QMessageBox.critical(self, "Ошибка", f"Произошла ошибка:\n{result_data['error']}")
            return

        self.last_anonymized_result = result_data
        self.current_file_replacements = result_data.get("file_replacements", {})
        
        user_entities = {alias: info for alias, info in self.current_mapping.items()
                        if info.get('category') == 'USER'}
        new_mapping = result_data.get("mapping", {})
        new_mapping.update(user_entities)
        
        self.current_mapping = new_mapping
        self.deanon_reverse_map = Deanon().load_map_from_data(self.current_mapping)
        self.current_map_name = 'только что созданная карта'
        self.current_map_path = None

        mapping = self.current_mapping
        mapping_list = []
        for alias, info in mapping.items():
            mapping_list.append({
                "pseudonym": alias,
                "original": info.get("original", ""),
                "category": info.get("category", ""),
                "file": ", ".join(info.get("files", []))
            })

        self.mapping_table.clearContents()
        self.mapping_table.setRowCount(len(mapping_list))

        for i, item in enumerate(mapping_list):
            self.mapping_table.setItem(i, 0, QTableWidgetItem(item["pseudonym"]))
            self.mapping_table.setItem(i, 1, QTableWidgetItem(item["original"]))
            self.mapping_table.setItem(i, 2, QTableWidgetItem(item["category"]))
            self.mapping_table.setItem(i, 3, QTableWidgetItem(item["file"]))

        self.refresh_map_list()
        if self.map_combo.count():
            self.map_combo.setCurrentIndex(0)

        if self.selected_files:
            current_item = self.file_list_widget.currentItem()
            preview_file = current_item.text() if current_item else self.selected_files[0]
            replacements = result_data.get("file_replacements", {}).get(preview_file, [])
            preview_text = self._extract_text_from_file(preview_file)
            if preview_text:
                selected_value = None
                if self.mapping_table.currentRow() >= 0:
                    selected_value = self.mapping_table.item(self.mapping_table.currentRow(), 1).text()
                self.preview_browser.setHtml(self._highlight_replacements_in_html(preview_text, replacements, selected_value=selected_value))

        QMessageBox.information(self, "Готово",
                                f"Обработано файлов: {result_data.get('files_processed', 0)}\n"
                                f"Найдено сущностей: {len(mapping_list)}")

    def save_mapping_json(self):
        if not self.current_mapping and self.mapping_table.rowCount() == 0:
            QMessageBox.warning(self, "Внимание", "Нет данных для сохранения.")
            return

        if not self.current_mapping:
            self.current_mapping = {}
            for i in range(self.mapping_table.rowCount()):
                alias = self.mapping_table.item(i, 0).text() if self.mapping_table.item(i, 0) else ""
                original = self.mapping_table.item(i, 1).text() if self.mapping_table.item(i, 1) else ""
                category = self.mapping_table.item(i, 2).text() if self.mapping_table.item(i, 2) else ""
                if alias:
                    self.current_mapping[alias] = {
                        "original": original,
                        "category": category,
                        "files": []
                    }

        file_path, _ = QFileDialog.getSaveFileName(
            self, "Сохранить карту", "", "JSON файл (*.json)"
        )
        if file_path:
            with open(file_path, 'w', encoding='utf-8') as f:
                json.dump(self.current_mapping, f, ensure_ascii=False, indent=2)
            QMessageBox.information(self, "Успех", f"Карта сохранена в: {file_path}")

    def export_to_excel(self):
        try:
            import openpyxl
        except ImportError:
            QMessageBox.warning(
                self, "Ошибка",
                "Модуль openpyxl не установлен. Установите: pip install openpyxl"
            )
            return

        rows = self.mapping_table.rowCount()
        if rows == 0:
            QMessageBox.warning(self, "Внимание", "Нет данных для экспорта.")
            return

        file_path, _ = QFileDialog.getSaveFileName(
            self, "Сохранить в Excel", "", "Excel файл (*.xlsx)"
        )
        if file_path:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Карта соответствий"
            headers = ["Псевдоним", "Исходное значение", "Категория"]
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.font = openpyxl.styles.Font(bold=True)

            for i in range(rows):
                row_num = i + 2
                ws.cell(row=row_num, column=1,
                        value=self.mapping_table.item(i, 0).text() if self.mapping_table.item(i, 0) else "")
                ws.cell(row=row_num, column=2,
                        value=self.mapping_table.item(i, 1).text() if self.mapping_table.item(i, 1) else "")
                ws.cell(row=row_num, column=3,
                        value=self.mapping_table.item(i, 2).text() if self.mapping_table.item(i, 2) else "")

            wb.save(file_path)
            QMessageBox.information(self, "Успех", f"Экспорт в Excel выполнен: {file_path}")

    def _get_default_deanon_map(self):
        if self.deanon_reverse_map:
            return self.deanon_reverse_map

        if self.current_mapping:
            return Deanon().load_map_from_data(self.current_mapping)

        if self.last_anonymized_result and isinstance(self.last_anonymized_result, dict):
            mapping = self.last_anonymized_result.get('mapping', {})
            if mapping:
                user_entities = {alias: info for alias, info in self.current_mapping.items()
                                if info.get('category') == 'USER'}
                mapping.update(user_entities)
                self.current_mapping = mapping
                self.deanon_reverse_map = Deanon().load_map_from_data(mapping)
                return self.deanon_reverse_map

        if self.map_combo.count() > 0:
            idx = self.map_combo.currentIndex()
            if idx >= 0:
                path = self.map_combo.itemData(idx)
                if path:
                    try:
                        mapping = self.map_manager.load_map(path)
                        user_entities = {alias: info for alias, info in self.current_mapping.items()
                                        if info.get('category') == 'USER'}
                        mapping.update(user_entities)
                        self.current_mapping = mapping
                        self.deanon_reverse_map = Deanon().load_map_from_data(mapping)
                        return self.deanon_reverse_map
                    except Exception:
                        pass

        return {}

    def deanonymize_window(self):
        """Открывает окно деанонимизации с поддержкой нескольких файлов и выделением сущностей"""
        deanon_window = DeanonWindow(
            parent=self,
            map_manager=self.map_manager,
            last_mapping=self.current_mapping or {},
            last_map_path=self.current_map_path,
            stylesheet_fn=self._build_stylesheet
        )
        deanon_window.exec_()

    def show_settings(self):
        dialog = SettingsDialog(self, initial_settings=self.settings.copy())
        if self.save_folder:
            folder_name = os.path.basename(self.save_folder)
            dialog.save_folder_btn.setText(f"Папка: {folder_name}")


        if dialog.exec_() == QDialog.Accepted:
            self.settings = dialog.get_settings()
            self.refresh_controls()
            if self.selected_files:
                self.show_file_preview(self.selected_files[0])
            QMessageBox.information(self, "Настройки",
                                    f"Сохранены настройки для категорий: {list(self.settings.keys())}")

# core/processors/replacer.py
import os
from docx import Document
from docx.shared import Pt
from PIL import Image, ImageDraw, ImageFont
from core.base_parser import InternalDocument, Paragraph as InternalParagraph


class Replacer:
    def __init__(self):
        self.font_path = "C:/Windows/Fonts/arial.ttf"

    @staticmethod
    def _filter_replacements_for_text(text: str, replacements: list) -> list:
        if not text or not replacements:
            return []

        filtered = []
        seen = set()
        lowered = text.lower()

        for rep in replacements:
            value = rep.get('value')
            alias = rep.get('alias')
            if value is None or alias is None:
                continue

            value_str = str(value).strip()
            alias_str = str(alias).strip()
            if not value_str or not alias_str:
                continue

            start = rep.get('start')
            end = rep.get('end')

            if isinstance(start, int) and isinstance(end, int) and 0 <= start < end <= len(text):
                fragment = text[start:end]
                if fragment.lower() == value_str.lower():
                    key = (start, end, alias_str)
                    if key in seen:
                        continue
                    seen.add(key)
                    filtered.append({
                        'start': start,
                        'end': end,
                        'value': value_str,
                        'alias': alias_str,
                    })
                    continue

            search_from = 0
            while True:
                idx = lowered.find(value_str.lower(), search_from)
                if idx == -1:
                    break
                end_pos = idx + len(value_str)
                key = (idx, end_pos, alias_str)
                if key not in seen:
                    seen.add(key)
                    filtered.append({
                        'start': idx,
                        'end': end_pos,
                        'value': value_str,
                        'alias': alias_str,
                    })
                search_from = idx + 1

        return sorted(filtered, key=lambda item: item['start'], reverse=True)

    @staticmethod
    def apply_replacements_to_text(text: str, replacements: list) -> str:
        if not text or not replacements:
            return text

        filtered = Replacer._filter_replacements_for_text(text, replacements)
        if not filtered:
            return text

        result = text
        for rep in filtered:
            start = int(rep['start'])
            end = int(rep['end'])
            alias = str(rep['alias'])
            if 0 <= start < end <= len(result):
                result = result[:start] + alias + result[end:]
            else:
                value = str(rep['value'])
                if value in result:
                    result = result.replace(value, alias)
        return result

    def build_docx(self, model: InternalDocument, replacements: list, output_path: str):
        doc = Document()

        for page in model.pages:
            for para in page.paragraphs:
                self._add_paragraph(doc, para, replacements)

            for table in page.tables:
                self._add_table(doc, table, replacements)

        doc.save(output_path)

    def _add_paragraph(self, doc, para: InternalParagraph, replacements: list):
        """Добавляет параграф в документ с заменой текста."""
        if not para.fragments:
            return
        new_para = doc.add_paragraph()
        combined_text = ''.join(frag.text for frag in para.fragments)
        block_replacements = self._filter_replacements_for_text(combined_text, replacements)
        text = self.apply_replacements_to_text(combined_text, block_replacements)
        run = new_para.add_run(text)
        for frag in para.fragments:
            if frag.bold:
                run.bold = True
                break
        for frag in para.fragments:
            if frag.italic:
                run.italic = True
                break
        for frag in para.fragments:
            if frag.underline:
                run.underline = True
                break
        if para.fragments and para.fragments[0].size:
            run.font.size = Pt(para.fragments[0].size)
        if para.fragments and para.fragments[0].font:
            run.font.name = para.fragments[0].font

    def _add_table(self, doc, table, replacements: list):
        """Добавляет таблицу из модели в документ с заменой текста."""
        if not table.rows:
            return

        num_rows = len(table.rows)
        num_cols = max(len(row) for row in table.rows) if table.rows else 0
        if num_cols == 0:
            return

        doc_table = doc.add_table(rows=num_rows, cols=num_cols)
        doc_table.style = 'Table Grid'

        for i, row in enumerate(table.rows):
            for j, cell in enumerate(row):
                if j >= num_cols:
                    continue
                doc_cell = doc_table.cell(i, j)
                doc_cell.paragraphs.clear()
                for para in cell.paragraphs:
                    if not para.fragments:
                        continue
                    cell_para = doc_cell.add_paragraph()
                    combined_text = ''.join(frag.text for frag in para.fragments)
                    block_replacements = self._filter_replacements_for_text(combined_text, replacements)
                    text = self.apply_replacements_to_text(combined_text, block_replacements)
                    run = cell_para.add_run(text)
                    if para.fragments and para.fragments[0].bold:
                        run.bold = True
                    if para.fragments and para.fragments[0].italic:
                        run.italic = True
                    if para.fragments and para.fragments[0].underline:
                        run.underline = True
                    if para.fragments and para.fragments[0].size:
                        run.font.size = Pt(para.fragments[0].size)
                    if para.fragments and para.fragments[0].font:
                        run.font.name = para.fragments[0].font

    def build_txt(self, model: InternalDocument, replacements: list, output_path: str):
        lines = []
        for page in model.pages:
            for para in page.paragraphs:
                text = ''.join(f.text for f in para.fragments)
                block_replacements = self._filter_replacements_for_text(text, replacements)
                text = self.apply_replacements_to_text(text, block_replacements)
                if text.strip():
                    lines.append(text)
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write('\n'.join(lines))

    def build_image(self, model: InternalDocument, replacements: list, output_path: str, original_image_path: str):
        img = Image.open(original_image_path).convert('RGB')
        draw = ImageDraw.Draw(img)
        try:
            font = ImageFont.truetype(self.font_path, 16)
        except Exception:
            font = ImageFont.load_default()
        for page in model.pages:
            for para in page.paragraphs:
                for frag in para.fragments:
                    if frag.bbox is None:
                        continue
                    x0, y0, x1, y1 = frag.bbox
                    new_text = self.apply_replacements_to_text(frag.text, replacements)
                    draw.rectangle([x0, y0, x1, y1], fill='white')
                    draw.text((x0, y0), new_text, fill='black', font=font)
        img.save(output_path)

    def build_pdf(self, model: InternalDocument, replacements: list, output_path: str, original_pdf_path: str = None):
        import fitz
        doc_new = fitz.open()
        doc_orig = None
        if original_pdf_path and os.path.exists(original_pdf_path):
            doc_orig = fitz.open(original_pdf_path)
        for page_idx, page in enumerate(model.pages):
            if page.width and page.height:
                pdf_page = doc_new.new_page(width=page.width, height=page.height)
            else:
                pdf_page = doc_new.new_page()
            if doc_orig and page_idx < len(doc_orig):
                orig_page = doc_orig[page_idx]
                pix = orig_page.get_pixmap()
                imgdata = pix.tobytes('png')
                pdf_page.insert_image(pdf_page.rect, stream=imgdata)
            for para in page.paragraphs:
                for frag in para.fragments:
                    if frag.bbox is None:
                        continue
                    x0, y0, x1, y1 = frag.bbox
                    rect = fitz.Rect(x0, y0, x1, y1)
                    pdf_page.draw_rect(rect, color=(1, 1, 1), fill=(1, 1, 1))
                    new_text = self.apply_replacements_to_text(frag.text, replacements)
                    fontsize = max(8, y1 - y0)
                    pdf_page.insert_text((x0, y0 + fontsize), new_text, fontsize=fontsize, fontname='helv')
        if doc_orig:
            doc_orig.close()
        doc_new.save(output_path)
        doc_new.close()

    def replace(self, model: InternalDocument, replacements: list, original_path: str, output_path: str,
                output_format: str = 'docx'):
        if output_format == 'docx':
            self.build_docx(model, replacements, output_path)
        elif output_format == 'txt':
            self.build_txt(model, replacements, output_path)
        elif output_format == 'image':
            self.build_image(model, replacements, output_path, original_path)
        elif output_format == 'pdf':
            self.build_pdf(model, replacements, output_path, original_path)
        else:
            raise ValueError(f'Неподдерживаемый формат: {output_format}')
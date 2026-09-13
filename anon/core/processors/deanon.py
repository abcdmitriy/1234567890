# core/processors/deanon.py
import json
import re
from docx import Document
from typing import Dict, Optional

class Deanon:
    def __init__(self):
        self.reverse_map = {}

    @staticmethod
    def _normalize_mapping(mapping_obj):
        if isinstance(mapping_obj, dict) and 'mapping' in mapping_obj and isinstance(mapping_obj['mapping'], dict):
            return mapping_obj['mapping']
        if isinstance(mapping_obj, dict):
            return mapping_obj
        if isinstance(mapping_obj, list):
            result = {}
            for item in mapping_obj:
                if isinstance(item, dict):
                    alias = item.get('pseudonym') or item.get('alias')
                    if alias:
                        result[alias] = {
                            'original': item.get('original', ''),
                            'category': item.get('category', ''),
                            'files': item.get('files', [])
                        }
            return result
        raise ValueError("Неверный формат карты соответствия")

    def load_map(self, map_path: str) -> Dict[str, str]:
        """
        Загружает карту соответствия из JSON-файла и строит обратный словарь.
        Возвращает обратный словарь {псевдоним: оригинал}.
        """
        with open(map_path, 'r', encoding='utf-8') as f:
            raw = json.load(f)
        mapping = self._normalize_mapping(raw)
        self.reverse_map = {}
        for alias, info in mapping.items():
            record = info if isinstance(info, dict) else {'original': info}
            original = record.get('original')
            if original is not None:
                self.reverse_map[str(alias)] = str(original)
        return self.reverse_map

    def load_map_from_data(self, mapping: Dict[str, Dict]) -> Dict[str, str]:
        """
        Принимает готовую карту (как из маппера) и строит обратный словарь.
        """
        normalized = self._normalize_mapping(mapping)
        self.reverse_map = {}
        for alias, info in normalized.items():
            record = info if isinstance(info, dict) else {'original': info}
            original = record.get('original')
            if original is not None:
                self.reverse_map[str(alias)] = str(original)
        return self.reverse_map

    def deanon_text(self, text: str, reverse_map: Optional[Dict[str, str]] = None) -> str:
        """
        Заменяет все псевдонимы в тексте на оригиналы.
        Если reverse_map не передан, используется загруженный ранее.
        """
        if reverse_map is None:
            reverse_map = self.reverse_map
        if not reverse_map:
            raise ValueError("Карта соответствия не загружена. Вызовите load_map() или передайте reverse_map.")

        # Сортируем псевдонимы по убыванию длины, чтобы сначала заменить длинные
        for alias in sorted(reverse_map.keys(), key=len, reverse=True):
            original = reverse_map[alias]
            text = text.replace(alias, original)
            text = text.replace(f"[{alias}]", original)
        return text

    def deanon_docx(self, input_path: str, output_path: str, reverse_map: Optional[Dict[str, str]] = None) -> None:
        """
        Открывает DOCX, заменяет все псевдонимы на оригиналы и сохраняет новый файл.
        """
        if reverse_map is None:
            reverse_map = self.reverse_map
        if not reverse_map:
            raise ValueError("Карта соответствия не загружена.")

        doc = Document(input_path)

        # Обработка всех параграфов (включая таблицы)
        self._replace_in_paragraphs(doc.paragraphs, reverse_map)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    self._replace_in_paragraphs(cell.paragraphs, reverse_map)

        # Обработка колонтитулов (если нужны)
        for section in doc.sections:
            self._replace_in_paragraphs(section.header.paragraphs, reverse_map)
            self._replace_in_paragraphs(section.footer.paragraphs, reverse_map)

        doc.save(output_path)

    def _replace_in_paragraphs(self, paragraphs, reverse_map: Dict[str, str]) -> None:
        """
        Заменяет псевдонимы на оригиналы в runs всех параграфов.
        """
        # Сортируем ключи по убыванию длины
        sorted_aliases = sorted(reverse_map.keys(), key=len, reverse=True)

        for para in paragraphs:
            for run in para.runs:
                if not run.text:
                    continue
                new_text = run.text
                for alias in sorted_aliases:
                    if alias in new_text:
                        new_text = new_text.replace(alias, reverse_map[alias])
                if new_text != run.text:
                    run.text = new_text